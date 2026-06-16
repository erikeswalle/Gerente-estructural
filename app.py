import streamlit as st
import os
import json
import pandas as pd
import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
import google.generativeai as genai
from dotenv import load_dotenv, set_key
import io
import shutil
import numpy as np
import zipfile
from PIL import Image, ImageOps
import pyrebase
import stripe

# Cargar variables de entorno
load_dotenv()

# Configuración de la página
st.set_page_config(page_title="Auditor QA/QC ITIFE", page_icon="🏗️", layout="wide")

# Título Principal
st.title("🏗️ Sistema de Auditoría Estructural y QA/QC")

# ==========================================
# MÓDULO DE AUTENTICACIÓN (FIREBASE)
# ==========================================
firebaseConfig = {
  "apiKey": "AIzaSyCkWdtHqfYEPjrHORGn7D3juCprD2COFxk",
  "authDomain": "gerente-estructural.firebaseapp.com",
  "projectId": "gerente-estructural",
  "storageBucket": "gerente-estructural.firebasestorage.app",
  "messagingSenderId": "354825313447",
  "appId": "1:354825313447:web:c9fd2a5b3c0021a83ccf82",
  "measurementId": "G-FCKH3TRK0S",
  "databaseURL": "https://gerente-estructural-default-rtdb.firebaseio.com/"
}

# Inicializar Firebase
firebase = pyrebase.initialize_app(firebaseConfig)
auth = firebase.auth()

if 'user' not in st.session_state:
    st.session_state['user'] = None

if st.session_state['user'] is None:
    st.subheader("🔐 Acceso Restringido")
    st.markdown("Por favor, inicia sesión o crea una cuenta para acceder a la herramienta.")
    
    col_login, col_empty = st.columns([1, 1])
    with col_login:
        auth_mode = st.radio("Elige una opción:", ["Iniciar Sesión", "Registrarse"], horizontal=True)
        email = st.text_input("Correo electrónico")
        password = st.text_input("Contraseña", type="password")
        
        if auth_mode == "Iniciar Sesión":
            if st.button("🚪 Entrar", use_container_width=True):
                if not email or not password:
                    st.warning("Ingresa correo y contraseña.")
                else:
                    with st.spinner("Autenticando..."):
                        try:
                            user = auth.sign_in_with_email_and_password(email, password)
                            st.session_state['user'] = user
                            st.rerun()
                        except Exception as e:
                            st.error("Credenciales inválidas o usuario no encontrado.")
        else:
            if st.button("📝 Crear Cuenta", use_container_width=True):
                if not email or not password:
                    st.warning("Ingresa correo y contraseña.")
                else:
                    with st.spinner("Creando cuenta..."):
                        try:
                            user = auth.create_user_with_email_and_password(email, password)
                            st.success("✅ Cuenta creada exitosamente. Ahora puedes Iniciar Sesión.")
                        except Exception as e:
                            st.error(f"Error al crear cuenta. La contraseña debe tener al menos 6 caracteres y el correo ser válido.")
    
    st.stop() # Detiene la ejecución del resto de la app hasta que haya login exitoso

# ==========================================
# APP PRINCIPAL (USUARIOS LOGUEADOS)
# ==========================================
st.sidebar.success(f"Logueado como: {st.session_state['user'].get('email', 'Usuario')}")
if st.sidebar.button("Cerrar Sesión"):
    st.session_state['user'] = None
    st.rerun()

# --- Integración Stripe y Base de Datos ---
stripe.api_key = os.getenv("STRIPE_SECRET_KEY", "")
db = firebase.database()
uid = st.session_state['user']['localId']

# 1. Verificar si viene de un pago exitoso (Redirección de Stripe)
if "session_id" in st.query_params:
    session_id = st.query_params["session_id"]
    try:
        checkout_session = stripe.checkout.Session.retrieve(session_id)
        if checkout_session.payment_status == "paid":
            # Actualizar DB: el usuario es premium
            db.child("users").child(uid).update({"is_premium": True})
            st.success("✅ ¡Pago confirmado! Gracias por tu suscripción.")
            st.balloons()
    except Exception as e:
        st.error("Error al verificar el pago con Stripe.")

# 2. Leer estado de suscripción del usuario desde Firebase
user_data = db.child("users").child(uid).get().val()
is_premium = False
if user_data and user_data.get("is_premium"):
    is_premium = True

# 3. Bloqueo (Muro de Pago)
if not is_premium:
    st.error("🔒 Tu cuenta no tiene una suscripción activa.")
    st.markdown("Para utilizar el motor de Inteligencia Artificial y procesar tus planos, necesitas una membresía Pro.")
    
    # Crear sesión de Stripe en tiempo real
    try:
        session = stripe.checkout.Session.create(
            payment_method_types=['card'],
            line_items=[{
                'price_data': {
                    'currency': 'mxn',
                    'product_data': {
                        'name': 'Gerente Estructural Pro (Mensual)',
                        'description': 'Acceso ilimitado al motor de revisión automatizada.'
                    },
                    'unit_amount': 49900, # $499.00 MXN
                    'recurring': {'interval': 'month'}
                },
                'quantity': 1,
            }],
            mode='subscription',
            success_url='https://app.aura-workspace.com/?session_id={CHECKOUT_SESSION_ID}',
            cancel_url='https://app.aura-workspace.com/',
            client_reference_id=uid
        )
        st.link_button("💳 Proceder al Pago Seguro ($499 MXN/mes)", session.url, type="primary")
    except Exception as e:
        st.error(f"No se pudo conectar con el banco: {e}")
        
    st.stop() # Bloquea el resto de la aplicación

# --- Resto de la aplicación desbloqueada ---
st.markdown("""
Esta herramienta cruza información de Planos Estructurales (PDF) y Catálogos/Generadores (Excel) 
para detectar discrepancias e incumplimientos normativos usando Inteligencia Artificial.
""")

# Sidebar para Configuración
st.sidebar.header("⚙️ Menú Principal")

with st.sidebar.expander("⚙️ Configuración Avanzada", expanded=False):
    saved_key = os.getenv("GEMINI_API_KEY", "")
    api_key = st.text_input("Ingresa tu Gemini API Key", value=saved_key, type="password")

    if api_key and api_key != saved_key:
        set_key(".env", "GEMINI_API_KEY", api_key)
        st.caption("✅ API Key guardada en automático para futuras sesiones.")
st.sidebar.markdown("---")
modulo_activo = st.sidebar.radio("Selecciona Módulo Activo:", ["🤖 Auditoría Integral (Planos + Catálogos)", "📐 Auditoría Exclusiva de Planos"])
st.sidebar.markdown("---")
st.sidebar.subheader("Información del Proyecto")
project_name = st.sidebar.text_input("Bautiza el Proyecto (Requerido)", placeholder="Ej. Almacén Adquisiciones")

st.sidebar.info("Modo Nube: Los archivos se procesarán en memoria.")

st.sidebar.markdown("---")
st.sidebar.subheader("⚙️ Parámetros del Proyecto")
sis_est = st.sidebar.selectbox("Sistema Estructural Principal", [
    "No especificado (Deducción libre)",
    "Marcos Rígidos (Acero / Concreto)",
    "Marcos Rígidos de Acero Estructural",
    "Marcos Rígidos de Concreto Reforzado",
    "Marcos Rígidos + Muros de Mampostería",
    "Muros de Carga (Mampostería / Concreto)",
    "Sistema Mixto o Dual",
    "Estructura Ligera / Naves Industriales"
])
dem_sis = st.sidebar.selectbox("Nivel de Demanda Sísmica", [
    "No especificado",
    "Alta Demanda (Zona Sísmica D, E, F)",
    "Demanda Moderada",
    "Baja Demanda (Zona no sísmica)"
])
clas_edif = st.sidebar.selectbox("Clasificación de Edificación", [
    "No especificada",
    "Grupo A (Esencial: Hospitales, Escuelas)",
    "Grupo B (Convencional: Vivienda, Oficinas)"
])
met_dis = st.sidebar.selectbox("Método de Diseño (Factores)", [
    "No especificado",
    "LRFD (Diseño por Factores de Carga y Resistencia)",
    "ASD (Diseño por Esfuerzos Permisibles)"
])

st.sidebar.markdown("---")
st.sidebar.subheader("🧠 Criterios de Revisión Activos")

# Leer archivo de reglas genéricas
reglas_path = "heuristicas_auditoria.txt"
reglas_activas = ""
if os.path.exists(reglas_path):
    with open(reglas_path, "r", encoding="utf-8") as f:
        reglas_activas = f.read()
    with st.sidebar.expander("Ver Documento de Reglas Activo"):
        st.text(reglas_activas)
else:
    st.sidebar.warning("Archivo de reglas no encontrado.")

# ==========================================
# MÓDULOS DE AUDITORÍA
# ==========================================
if modulo_activo in ["🤖 Auditoría Integral (Planos + Catálogos)", "📐 Auditoría Exclusiva de Planos"]:
    is_integral = modulo_activo == "🤖 Auditoría Integral (Planos + Catálogos)"
    st.subheader(f"{modulo_activo.split(' ')[0]} Módulo de {modulo_activo.split(' ', 1)[1]}")
    
    # Área de Carga de Archivos
    if is_integral:
        col1, col2 = st.columns(2)
        with col1:
            st.subheader("📄 Carga de Planos (PDF)")
            pdf_files = st.file_uploader("Sube los planos estructurales", type=["pdf"], accept_multiple_files=True)
        with col2:
            st.subheader("📊 Carga de Cantidades (Excel)")
            excel_files = st.file_uploader("Sube el catálogo de conceptos y/o generadores", type=["xlsx"], accept_multiple_files=True)
    else:
        st.subheader("📄 Carga de Planos (PDF)")
        pdf_files = st.file_uploader("Sube los planos estructurales", type=["pdf"], accept_multiple_files=True)
        excel_files = []
    
    st.markdown("---")

# --- FUNCIONES CORE DEL MOTOR ---

def extract_pdf_text(pdf_upload):
    doc = fitz.open(stream=pdf_upload.read(), filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text()
    pdf_upload.seek(0)
    return text

def extract_excel_data(excel_upload):
    df_dict = pd.read_excel(excel_upload, sheet_name=None)
    excel_upload.seek(0)
    
    summary = ""
    for sheet_name, df in df_dict.items():
        summary += f"--- HOJA: {sheet_name} ---\n"
        # Convertir a string las primeras filas para contexto
        summary += df.head(50).to_csv(index=False) + "\n\n"
    return summary

def call_gemini(api_key, pdf_texts, excel_texts, reglas, tipo_auditoria="integral", sis_est="", dem_sis="", clas_edif="", met_dis=""):
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel('gemini-3.5-flash')
    
    if tipo_auditoria == "integral":
        excel_context = f"\nDATOS EXTRAÍDOS DE EXCEL:\n{excel_texts}" if excel_texts.strip() else ""
        mandato_extra = " y Catálogos/Generadores (Excels)."
        regla_excel = " Elementos dibujados deben coincidir estrictamente con lo cobrado en los Excels."
    else:
        excel_context = ""
        mandato_extra = " EXCLUSIVAMENTE (No hay excels en esta corrida)."
        regla_excel = ""

    prompt = f"""
    [IDENTIDAD Y ROL]
    Actúa como un Gerente de Proyectos Estructurales nivel Senior y Director Responsable de Obra (DRO) con más de 20 años de experiencia comprobable. Tu firma avala la seguridad estructural, viabilidad constructiva y cumplimiento legal del proyecto. No eres un simple revisor de dibujo; eres el último filtro técnico antes de la construcción.

    [MANDATO PRINCIPAL]
    Tu tarea es auditar de manera implacable e hiper-estricta la información de los siguientes textos extraídos de Planos (PDFs){mandato_extra} Tu revisión no se basa en heurísticas genéricas, sino en la aplicación dogmática de las normativas vigentes internacionales y locales seleccionadas.

    [CÓDIGOS Y NORMAS A APLICAR]
    {reglas}

    [DIRECTRICES DE AUDITORÍA AVANZADA]
    1. Ingeniería de Detalle en Acero Estructural: Exige el máximo rigor. Si detectas conexiones, audita la viabilidad de placas base, pernos/anclas, atiesadores y soldaduras. Si un detalle falla por cálculo, EXPLICA CLARAMENTE por qué no pasa, y PROPÓN ESPECÍFICAMENTE la solución completa: qué tipo/grado de perno usar, separación requerida, espesor/tamaño de placa y/o especificación de soldadura. Si el error es una "incongruencia de dibujo" (ej. usar un mismo detalle para dos soluciones distintas), indícale al calculista que debe separar y crear detalles independientes.
    2. Ingeniería de Detalle en Concreto Reforzado: Audita que el diseño garantice la ductilidad y resistencia. Revisa estrictamente zonas de confinamiento sísmico (separación de estribos en nudos y extremos), longitudes de desarrollo, traslapes y cuantías mínimas de acero según el ACI 318 / NTC.
    3. Congruencia Geométrica Espacial: Realiza un escaneo cruzado implacable de la topografía y geometría. Las cotas, los ejes y los niveles de elevación (N.P.T., N.L.B., etc.) deben cuadrar perfectamente al milímetro en todos los planos (plantas, cortes, fachadas y detalles). Levanta alerta crítica ante cualquier incongruencia espacial.
    4. Trazabilidad de Materiales: Las resistencias a compresión (f'c), límite de fluencia del acero (fy) y módulos de elasticidad especificados en los planos deben ser congruentes en el 100% de la documentación.{regla_excel}
    5. Higiene Documental y Nomenclatura: Castiga los errores de "Copy-Paste" del proyectista. Usa tu inteligencia como estructurista: evalúa si existe NOMENCLATURA REPETIDA (ej. usar la misma etiqueta 'PB-1' para detallar dos anclajes con dimensiones/armados FÍSICAMENTE INCOMPATIBLES O CONTRADICTORIOS). Si la misma etiqueta se usa para dos detalles distintos, márcalo como "Nomenclatura repetida/incongruente". (Ojo: si es la misma placa repetida en varios lugares con las mismas propiedades, NO es error).

    [CONTEXTO DEL PROYECTO]
    Sistema Estructural Principal: {sis_est}
    Nivel de Demanda Sísmica: {dem_sis}
    Clasificación de Edificación (Importancia): {clas_edif}
    Método de Diseño (Factores de Carga): {met_dis}

    [DATOS EXTRAÍDOS]
    {pdf_texts}
    {excel_context}

    [FORMATO DE RESPUESTA Y TONO]
    Tu tono debe ser el de un perito estructural: directo, técnico, sin rodeos y enfocado en el riesgo y la solución normativa.
    Claridad Ejecutiva: Explica los errores en la columna "error" de forma que una persona común lo entienda (Distingue claramente si es un error de "Congruencia de Dibujo / Nomenclatura Repetida" o una "Falla Estructural Matemática"), pero respalda tu rechazo sin perder el rigor ingenieril en la memoria de cálculo.
    Análisis de Comportamiento: Antes de emitir una observación matemática, explica brevemente en la Memoria de Cálculo el COMPORTAMIENTO FÍSICO del sistema elegido ({sis_est}) y por qué el error es crítico para su zona sísmica y tipo de edificio.
    Obligatorio: Toda observación debe estar FUNDAMENTADA, citando (en la medida de lo posible) el código, sección o manual específico que se está incumpliendo (ej. ACI 318, AISC 360, AWS D1.1, NTC).
    Unidades (SI): Utiliza estrictamente el Sistema Internacional (SI) y métrico decimal en todos los cálculos y propuestas (ej. kgf, toneladas, cm², kgf/m, mm). Evita usar kips, libras, o MPa sin su equivalente de la práctica común local.
    Usa saltos de línea (\\n) dentro de tus textos para separar párrafos o viñetas y mantener todo muy legible.
    
    DEBES RESPONDER ÚNICAMENTE CON UN JSON VÁLIDO CON ESTA ESTRUCTURA EXACTA (sin markdown adicional):
    [
        {{
            "tipo_archivo": "PDF",
            "archivo_nombre": "nombre_del_pdf_detectado",
            "error": "[Cita Normativa, ej. AISC 360-16 J2] - Descripción del incumplimiento - Propuesta de solución.",
            "memoria_calculo": "Demuestra matemáticamente por qué el detalle original FALLA (ej. demanda vs capacidad) y justifica con cálculos por qué TU PROPUESTA SÍ PASA. Usa saltos de línea (\\n) para separar operaciones.",
            "severidad": "Alta/Media/Baja",
            "texto_a_buscar": "Texto corto y exacto que aparece en el PDF donde está el error para enmarcarlo"
        }},
        {{
            "tipo_archivo": "EXCEL",
            "archivo_nombre": "nombre_del_excel_detectado",
            "error": "Descripción detallada de la omisión volumétrica o incongruencia",
            "memoria_calculo": "Desglose matemático de las cantidades analizadas o diferencias volumétricas.",
            "severidad": "Alta/Media/Baja",
            "texto_a_buscar": "Valor exacto en la celda de Excel para marcarla"
        }}
    ]
    Si no hay errores, responde [].
    """
    
    response = model.generate_content(prompt)
    try:
        # Limpiar posible markdown ```json
        raw_text = response.text.replace("```json", "").replace("```", "").strip()
        return json.loads(raw_text)
    except Exception as e:
        st.error(f"Error parseando JSON de Gemini: {e}")
        return []

def generate_advanced_excel_report(gemini_results, pdf_files):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, Alignment
    
    wb = openpyxl.Workbook()
    
    # 1. Hoja "Índice de Planos"
    ws_index = wb.active
    ws_index.title = "Índice de Planos"
    ws_index.append(["Nombre del Plano", "Total de Observaciones"])
    for cell in ws_index[1]:
        cell.font = Font(bold=True)
    
    conteo_planos = {}
    if isinstance(gemini_results, list):
        for issue in gemini_results:
            if issue.get("tipo_archivo") == "PDF":
                plano = issue.get("archivo_nombre", "Desconocido")
                conteo_planos[plano] = conteo_planos.get(plano, 0) + 1
            
    nombres_subidos = [pdf.name for pdf in pdf_files] if pdf_files else []
    for nombre in nombres_subidos:
        ws_index.append([nombre, conteo_planos.get(nombre, 0)])
        
    ws_index.column_dimensions['A'].width = 50
    ws_index.column_dimensions['B'].width = 25
    
    # 2. Hoja "Reporte Fotográfico"
    ws_obs = wb.create_sheet(title="Reporte Fotográfico")
    headers = ["Plano", "Severidad", "Observación", "Texto Detectado", "Memoria de Cálculo", "Evidencia Visual"]
    ws_obs.append(headers)
    for cell in ws_obs[1]:
        cell.font = Font(bold=True)
        
    ws_obs.column_dimensions['A'].width = 30
    ws_obs.column_dimensions['B'].width = 15
    ws_obs.column_dimensions['C'].width = 70
    ws_obs.column_dimensions['D'].width = 30
    ws_obs.column_dimensions['E'].width = 70 # Para Memoria de Calculo
    ws_obs.column_dimensions['F'].width = 50 # Para imagen
    
    pdf_dict = {}
    if pdf_files:
        for pdf in pdf_files:
            pdf_dict[pdf.name] = pdf
            
    row_num = 2
    if isinstance(gemini_results, list):
        for issue in gemini_results:
            plano = issue.get("archivo_nombre", "Desconocido")
            sev = issue.get("severidad", "")
            obs = issue.get("error", "")
            texto = issue.get("texto_a_buscar", "")
            memoria = issue.get("memoria_calculo", "")
            
            # Altura dinámica de la fila basada en el largo del texto para evitar que se oculte
            max_len = max(len(str(obs)), len(str(memoria)))
            estimated_height = max(30.0, (max_len / 70.0) * 15.0)
            ws_obs.row_dimensions[row_num].height = estimated_height
            
            ws_obs.cell(row=row_num, column=1, value=plano).alignment = Alignment(vertical='top', horizontal='left')
            ws_obs.cell(row=row_num, column=2, value=sev).alignment = Alignment(vertical='top', horizontal='left')
            ws_obs.cell(row=row_num, column=3, value=obs).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
            ws_obs.cell(row=row_num, column=4, value=texto).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
            ws_obs.cell(row=row_num, column=5, value=memoria).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
            
            # Intentar tomar screenshot si es PDF
            if issue.get("tipo_archivo") == "PDF" and plano in pdf_dict and texto:
                pdf_upload = pdf_dict[plano]
                try:
                    pdf_upload.seek(0)
                    doc = fitz.open(stream=pdf_upload.read(), filetype="pdf")
                    pdf_upload.seek(0)
                    img_inserted = False
                    
                    for page in doc:
                        if img_inserted: break
                        insts = page.search_for(texto)
                        if insts:
                            rect = insts[0]
                            # Agregar 60px de margen en cada dirección para dar contexto visual
                            clip_rect = rect + fitz.Rect(-60, -60, 60, 60)
                            pix = page.get_pixmap(clip=clip_rect, dpi=150)
                            
                            img_pil = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                            img_pil.thumbnail((260, 260))
                            img_pil = ImageOps.expand(img_pil, border=3, fill='black')
                            img_byte_arr = io.BytesIO()
                            img_pil.save(img_byte_arr, format='PNG')
                            img_byte_arr.seek(0)
                            
                            xl_img = XLImage(img_byte_arr)
                            current_h = ws_obs.row_dimensions[row_num].height or 15.0
                            ws_obs.row_dimensions[row_num].height = max(200.0, current_h) # Adjust row height para imagen
                            
                            # Compensar posicionamiento de imagen en celda
                            ws_obs.add_image(xl_img, f"F{row_num}")
                            img_inserted = True
                            
                except Exception as e:
                    pass
                    
            row_num += 1
            
    out_excel = io.BytesIO()
    wb.save(out_excel)
    return out_excel.getvalue()

def mark_pdf(pdf_upload, results):
    pdf_upload.seek(0)
    doc = fitz.open(stream=pdf_upload.read(), filetype="pdf")
    pdf_upload.seek(0)
    
    for issue in results:
        if issue.get("tipo_archivo") == "PDF":
            texto_buscar = issue.get("texto_a_buscar", "")
            if texto_buscar:
                for page in doc:
                    text_instances = page.search_for(texto_buscar)
                    for inst in text_instances:
                        # Dibujar rectángulo rojo
                        page.draw_rect(inst, color=(1, 0, 0), width=1.5)
                        # Añadir anotación adhesiva
                        annot = page.add_text_annot(inst.tl, issue.get("error", "Error detectado"))
                        annot.set_colors(stroke=(1, 1, 0)) # Amarillo
                        annot.update()
                        
    out_pdf = io.BytesIO()
    doc.save(out_pdf)
    return out_pdf.getvalue()

def mark_excel(excel_upload, results):
    wb = openpyxl.load_workbook(excel_upload)
    excel_upload.seek(0)
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    
    for issue in results:
        if issue.get("tipo_archivo") == "EXCEL":
            texto_buscar = issue.get("texto_a_buscar", "")
            if texto_buscar:
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value and str(texto_buscar).lower() in str(cell.value).lower():
                                cell.fill = red_fill
                                comment = Comment(issue.get("error", "Error de Auditoría"), "QA/QC Bot")
                                cell.comment = comment
                                
    out_excel = io.BytesIO()
    wb.save(out_excel)
    return out_excel.getvalue()

# --- EJECUCIÓN ---

if modulo_activo in ["🤖 Auditoría Integral (Planos + Catálogos)", "📐 Auditoría Exclusiva de Planos"]:
    tipo_auditoria = "integral" if modulo_activo == "🤖 Auditoría Integral (Planos + Catálogos)" else "exclusiva"
    btn_text = "🚀 Iniciar Auditoría Integral" if tipo_auditoria == "integral" else "🚀 Iniciar Auditoría de Planos"
    
    if st.button(btn_text, use_container_width=True):
        if not api_key:
            st.error("⚠️ Por favor ingresa tu API Key de Gemini en el panel lateral.")
        elif not project_name:
            st.error("⚠️ Por favor bautiza el proyecto en el panel lateral antes de continuar.")
        elif not pdf_files and not excel_files:
            st.warning("⚠️ Sube al menos un archivo para analizar.")
        else:
            with st.spinner("🧠 Extrayendo datos y consultando a Gemini API... Esto puede tomar unos minutos."):
                
                # 1. Extracción
                all_pdf_texts = ""
                for pdf in pdf_files:
                    all_pdf_texts += f"--- ARCHIVO: {pdf.name} ---\n"
                    all_pdf_texts += extract_pdf_text(pdf) + "\n\n"
                    
                all_excel_texts = ""
                for exc in excel_files:
                    all_excel_texts += f"--- ARCHIVO: {exc.name} ---\n"
                    all_excel_texts += extract_excel_data(exc) + "\n\n"
                    
                # 2. IA Auditoría
                gemini_results = call_gemini(api_key, all_pdf_texts, all_excel_texts, reglas_activas, tipo_auditoria, sis_est, dem_sis, clas_edif, met_dis)
                
                if isinstance(gemini_results, list):
                    # Ordenar resultados alfabéticamente por archivo_nombre
                    gemini_results.sort(key=lambda x: str(x.get("archivo_nombre", "")))
                
                if not gemini_results:
                    st.success("✅ No se detectaron errores estructurales graves según las reglas proporcionadas.")
                else:
                    st.error(f"🚨 Se encontraron {len(gemini_results)} inconsistencias (Red Flags).")
                    
                    # Mostrar en pantalla
                    df_report = pd.DataFrame(gemini_results)
                    st.dataframe(df_report)
                    
                    # 3. Generar Salidas (Marcado) y Guardar en Memoria (Cloud-Ready)
                    st.subheader("📥 Sistema de Archivos y Descarga")
                    
                    zip_buffer = io.BytesIO()
                    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                        # Guardar Reporte Maestro Avanzado con Imágenes
                        try:
                            excel_bytes = generate_advanced_excel_report(gemini_results, pdf_files)
                            zip_file.writestr(f"Reporte_Maestro_{project_name}.xlsx", excel_bytes)
                        except Exception as e:
                            st.error(f"Error generando reporte fotográfico: {e}")
                            
                        # Guardar PDFs marcados
                        if pdf_files:
                            for pdf in pdf_files:
                                marked_pdf_bytes = mark_pdf(pdf, gemini_results)
                                zip_file.writestr(f"[REVISADO]_{pdf.name}", marked_pdf_bytes)
                                    
                        # Guardar Excels marcados
                        if excel_files:
                            for exc in excel_files:
                                marked_excel_bytes = mark_excel(exc, gemini_results)
                                zip_file.writestr(f"[REVISADO]_{exc.name}", marked_excel_bytes)
                    
                    # Notificar al usuario
                    st.success("✅ Documentos procesados y empaquetados exitosamente.")
                    
                    # Botón único de descarga ZIP
                    st.download_button(
                        label=f"📦 Descargar Proyecto Completo (.zip)",
                        data=zip_buffer.getvalue(),
                        file_name=f"Auditoria_{project_name}.zip",
                        mime="application/zip",
                        use_container_width=True
                    )

st.markdown("---")
st.caption("Desarrollado como Sistema Automatizado de Control de Calidad Estructural.")
